"""DFAT workbook adapter, with aliases grouped by the official reference number."""

import re
from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook

from trust_signal.connectors.base import SanctionsListing, SanctionsSnapshot
from trust_signal.connectors.http_source import HttpSource, SourceSchemaError


class AustraliaSanctionsAdapter(HttpSource):
    source_id = "au_dfat_sanctions"
    url = "https://www.dfat.gov.au/sites/default/files/Australian_Sanctions_Consolidated_List.xlsx"
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def fetch_snapshot(self) -> SanctionsSnapshot:
        response = self.get(self.url)
        try:
            with ZipFile(BytesIO(response.content)) as archive:
                if sum(member.file_size for member in archive.infolist()) > 128 * 1024 * 1024:
                    raise ValueError("Workbook expansion exceeds limit")
            workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
            try:
                rows = iter(workbook.worksheets[0].iter_rows(values_only=True))
                headers = None
                for _ in range(20):
                    candidate = [str(value or "").strip() for value in next(rows)]
                    if {"Reference", "Type", "Name of Individual or Entity", "Name Type"} <= set(candidate):
                        headers = candidate
                        break
                if headers is None:
                    raise ValueError("Required DFAT columns missing")
                grouped = {}
                for values in rows:
                    if not any(value is not None for value in values):
                        continue
                    row = dict(zip(headers, values, strict=True))
                    reference = str(row["Reference"] or "").strip()
                    match = re.fullmatch(r"([0-9]+)([a-z]*)", reference, re.IGNORECASE)
                    name = str(row["Name of Individual or Entity"] or "").strip()
                    if not match or not name:
                        raise ValueError("Invalid DFAT identity")
                    group = grouped.setdefault(match[1], {"primary": None, "aliases": []})
                    if str(row["Name Type"]).strip().casefold() == "primary name":
                        if group["primary"] is not None:
                            raise ValueError("Duplicate DFAT primary reference")
                        group["primary"] = row
                    else:
                        group["aliases"].append(name)
                listings = []
                for reference, group in grouped.items():
                    row = group["primary"]
                    if row is None:
                        raise ValueError("Alias has no primary listing")
                    kind = str(row["Type"]).strip().casefold()
                    if kind not in {"person", "individual", "entity", "vessel"}:
                        raise ValueError("Unknown DFAT subject type")
                    measures = []
                    for column in ("Targeted Financial Sanction", "Travel Ban", "Arms Embargo", "Maritime Restriction"):
                        flag = row.get(column)
                        if flag is not None and str(flag).strip().casefold() not in {"true", "false", ""}:
                            raise ValueError("Invalid DFAT measure flag")
                        if str(flag).strip().casefold() == "true":
                            measures.append(column)
                    listings.append(SanctionsListing(source_id=self.source_id, source_record_id=reference,
                        legal_name=str(row["Name of Individual or Entity"]).strip(),
                        list_type="individual" if kind == "person" else kind,
                        aliases=group["aliases"], measures=measures,
                        sanctions_regime=str(row.get("Committees") or "") or None,
                        comments=str(row.get("Listing Information") or "") or None))
                if not listings:
                    raise ValueError("Empty DFAT workbook")
            finally:
                workbook.close()
        except Exception as exc:
            raise SourceSchemaError("DFAT workbook failed validation.") from exc
        return SanctionsSnapshot(observation=self.observation(response, "dfat-consolidated.xlsx",
            {"record_count": len(listings)}, self.mime, binary=True), listings=listings)
